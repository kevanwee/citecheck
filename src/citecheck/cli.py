"""CLI. Reads .md/.txt (and .docx with the optional extra), prints a report, exits by policy."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from .extract import extract_citations
from .report import build_report
from .resolve import CompositeResolver, HttpExistenceResolver, ManifestResolver


def read_text(path: Path) -> str:
    if path.suffix.lower() == ".docx":
        try:
            import docx  # type: ignore
        except ImportError:
            sys.exit("reading .docx requires `pip install citecheck[docx]`")
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.extend(c.text for c in row.cells)
        # footnotes are where citations live in many drafts; python-docx does not expose
        # them, so pull the raw XML part as a fallback
        try:
            fn = d.part.package.part_related_by(
                "http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes"
            )
            import re
            parts.append(re.sub(r"<[^>]+>", " ", fn.blob.decode("utf-8", "ignore")))
        except (KeyError, AttributeError):
            pass
        return "\n".join(parts)
    return path.read_text(encoding="utf-8")


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(prog="citecheck",
                                description="Verify that every authority cited in a draft exists.")
    p.add_argument("draft", type=Path)
    p.add_argument("--manifest", "-m", type=Path, action="append", default=[],
                   help="JSON manifest of known authorities (repeatable)")
    p.add_argument("--online", action="store_true",
                   help="also check canonical URLs over HTTP (opt-in; honour site terms)")
    p.add_argument("--strict", action="store_true",
                   help="fail if any citation is unchecked or ambiguous")
    p.add_argument("--json", action="store_true")
    p.add_argument("--extract-only", action="store_true", help="list citations, no resolution")
    args = p.parse_args(argv)

    text = read_text(args.draft)
    citations = extract_citations(text)

    if args.extract_only:
        for c in citations:
            print(f"{c.key}\t{c.kind.value}\t{c.jurisdiction}\tx{c.occurrences}")
        return 0

    resolvers = [ManifestResolver(m) for m in args.manifest]
    if args.online:
        resolvers.append(HttpExistenceResolver())
    report = build_report(citations, CompositeResolver(resolvers), strict=args.strict)

    print(json.dumps(report.to_dict(), indent=2) if args.json else report.to_markdown())
    return report.exit_code


if __name__ == "__main__":
    sys.exit(main())
